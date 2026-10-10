"""Extracts clean text and page metadata from PDF, DOCX, TXT, and Markdown files."""
import os
from pathlib import Path
from typing import List, Dict, Any, Tuple


class TextExtractor:
    @staticmethod
    def extract_file(file_path: Path, file_type: str) -> Tuple[List[Dict[str, Any]], int]:
        """
        Extract text page-by-page or block-by-block.
        Returns:
            pages: List of {"page_number": int, "text": str}
            total_pages: int
        """
        file_type = file_type.lower().lstrip(".")
        if file_type == "pdf":
            return TextExtractor._extract_pdf(file_path)
        elif file_type == "docx":
            return TextExtractor._extract_docx(file_path)
        elif file_type in ["txt", "md"]:
            return TextExtractor._extract_plaintext(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    @staticmethod
    def _extract_pdf(file_path: Path) -> Tuple[List[Dict[str, Any]], int]:
        import fitz  # PyMuPDF
        pages = []
        doc = fitz.open(file_path)
        total_pages = len(doc)
        for page_idx in range(total_pages):
            page = doc[page_idx]
            text = page.get_text("text").strip()
            pages.append({
                "page_number": page_idx + 1,
                "text": text
            })
        doc.close()
        return pages, total_pages

    @staticmethod
    def _extract_docx(file_path: Path) -> Tuple[List[Dict[str, Any]], int]:
        import docx
        doc = docx.Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    full_text.append(row_text)
        
        joined = "\n\n".join(full_text)
        return [{"page_number": 1, "text": joined}], 1

    @staticmethod
    def _extract_plaintext(file_path: Path) -> Tuple[List[Dict[str, Any]], int]:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read().strip()
        return [{"page_number": 1, "text": content}], 1


class DocumentChunker:
    @staticmethod
    def chunk_pages(
        pages: List[Dict[str, Any]],
        chunk_size: int = 500,
        chunk_overlap: int = 80
    ) -> List[Dict[str, Any]]:
        """
        Chunks page text while preserving page number references and character boundaries.
        Returns:
            chunks: List of {"chunk_index": int, "content": str, "page_number": int, "token_count": int}
        """
        chunks = []
        chunk_index = 0

        for page in pages:
            text = page["text"]
            page_num = page["page_number"]
            if not text:
                continue

            # Split into words / simple tokens
            words = text.split()
            if not words:
                continue

            # If text is smaller than chunk size, emit as single chunk
            if len(words) <= chunk_size:
                chunks.append({
                    "chunk_index": chunk_index,
                    "content": text,
                    "page_number": page_num,
                    "token_count": len(words)
                })
                chunk_index += 1
                continue

            # Sliding window with overlap
            start = 0
            while start < len(words):
                end = min(start + chunk_size, len(words))
                chunk_words = words[start:end]
                chunk_text = " ".join(chunk_words)

                chunks.append({
                    "chunk_index": chunk_index,
                    "content": chunk_text,
                    "page_number": page_num,
                    "token_count": len(chunk_words)
                })
                chunk_index += 1

                if end == len(words):
                    break
                start += (chunk_size - chunk_overlap)

        return chunks
